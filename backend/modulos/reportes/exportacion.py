"""Exportación de reportes a Excel (openpyxl) y PDF (reportlab), HU-16.

Cada reporte se arma como una lista de `Seccion` (título, columnas, filas
y notas) y este módulo solo sabe dibujar secciones: no conoce kardex,
conciliación ni nada del negocio. Así un reporte nuevo no toca este
archivo.
"""
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from xml.sax.saxutils import escape


@dataclass
class Seccion:
    titulo: str
    columnas: List[str]
    filas: List[List[Any]]
    notas: List[str] = field(default_factory=list)


@dataclass
class DocumentoReporte:
    titulo: str
    encabezado: List[str]  # líneas: contribuyente, año gravable, fecha...
    secciones: List[Seccion]


FORMATO_NUMERO_EXCEL = "#,##0.00"
_TAMANO_CELDA = 7.5
_COLOR_ENCABEZADO = "1F4E79"


def _numero_co(valor: float) -> str:
    """1234567.5 -> '1.234.567,50' (formato colombiano)."""
    return f"{valor:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def _texto_celda(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, int):
        return str(valor)  # ids y conteos, sin decimales
    if isinstance(valor, float):
        return _numero_co(valor)
    if isinstance(valor, (date, datetime)):
        return valor.strftime("%Y-%m-%d")
    return str(valor)


# ------------------------------------------------------------------------ Excel


def a_excel(documento: DocumentoReporte) -> bytes:
    """Una sola hoja con las secciones apiladas, cada una con su título,
    encabezado de columnas y notas."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte"
    anchos: dict = {}

    hoja.append([documento.titulo])
    hoja.cell(row=hoja.max_row, column=1).font = Font(bold=True, size=14)
    for linea in documento.encabezado:
        hoja.append([linea])

    for seccion in documento.secciones:
        hoja.append([])
        hoja.append([seccion.titulo])
        hoja.cell(row=hoja.max_row, column=1).font = Font(bold=True, size=12)

        hoja.append(seccion.columnas)
        for i in range(1, len(seccion.columnas) + 1):
            celda = hoja.cell(row=hoja.max_row, column=i)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=_COLOR_ENCABEZADO)
            celda.alignment = Alignment(wrap_text=True, vertical="center")
            anchos[i] = max(anchos.get(i, 0), len(str(celda.value)))

        for fila in seccion.filas:
            hoja.append([_valor_excel(v) for v in fila])
            for i, valor in enumerate(fila, start=1):
                celda = hoja.cell(row=hoja.max_row, column=i)
                if isinstance(valor, float):
                    celda.number_format = FORMATO_NUMERO_EXCEL
                elif isinstance(valor, (date, datetime)):
                    celda.number_format = "yyyy-mm-dd"
                anchos[i] = max(anchos.get(i, 0), len(_texto_celda(valor)))

        for nota in seccion.notas:
            hoja.append([nota])
            hoja.cell(row=hoja.max_row, column=1).font = Font(italic=True)

    for i, ancho in anchos.items():
        hoja.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 10), 60)

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _valor_excel(valor: Any) -> Any:
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if hasattr(valor, "value"):  # Enum
        return valor.value
    return valor


# -------------------------------------------------------------------------- PDF


def a_pdf(documento: DocumentoReporte) -> bytes:
    """A4 horizontal; cada sección es una tabla que repite su encabezado si
    salta de página. Los textos largos se ajustan dentro de la celda."""
    buffer = io.BytesIO()
    pagina = landscape(A4)
    margen = 1.5 * cm
    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagina,
        leftMargin=margen,
        rightMargin=margen,
        topMargin=margen,
        bottomMargin=margen,
        title=documento.titulo,
    )
    estilos = getSampleStyleSheet()
    celda = estilos["BodyText"].clone("celda", fontSize=_TAMANO_CELDA, leading=9)
    celda_numero = celda.clone("celda_numero", alignment=TA_RIGHT)
    celda_encabezado = celda.clone("celda_encabezado", textColor=colors.white, fontName="Helvetica-Bold")
    nota = estilos["BodyText"].clone("nota", fontSize=8, textColor=colors.grey, fontName="Helvetica-Oblique")
    ancho_util = pagina[0] - 2 * margen

    elementos: list = [Paragraph(escape(documento.titulo), estilos["Title"])]
    for linea in documento.encabezado:
        elementos.append(Paragraph(escape(linea), estilos["Normal"]))

    for seccion in documento.secciones:
        elementos.append(Spacer(1, 0.4 * cm))
        elementos.append(Paragraph(escape(seccion.titulo), estilos["Heading3"]))

        if seccion.filas:
            textos = [[_texto_celda(_valor_excel(v)) for v in fila] for fila in seccion.filas]
            datos = [[Paragraph(escape(c), celda_encabezado) for c in seccion.columnas]]
            datos += [
                [
                    Paragraph(escape(t), celda_numero if _es_numero(v) else celda)
                    for t, v in zip(textos_fila, fila)
                ]
                for textos_fila, fila in zip(textos, seccion.filas)
            ]
            tabla = Table(
                datos,
                colWidths=_anchos_columnas(seccion.columnas, textos, ancho_util),
                repeatRows=1,
            )
            tabla.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_COLOR_ENCABEZADO}")),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
                    ]
                )
            )
            elementos.append(tabla)

        for texto in seccion.notas:
            elementos.append(Paragraph(escape(texto), nota))

    doc.build(elementos)
    return buffer.getvalue()


def _es_numero(valor: Any) -> bool:
    return isinstance(valor, (int, float)) and not isinstance(valor, bool)


def _anchos_columnas(columnas: List[str], filas: List[List[str]], ancho_util: float) -> List[float]:
    """Reparte el ancho de la página: cada columna recibe primero lo que
    ocupa su palabra más larga (para que ninguna palabra se parta) y el
    resto se reparte según el contenido más largo de cada una, con tope
    para que un texto largo no se coma la tabla. Medidas en puntos."""
    relleno = 14  # padding izquierdo + derecho de la celda, con holgura

    def ancho(texto: str, negrita: bool = False) -> float:
        fuente = "Helvetica-Bold" if negrita else "Helvetica"
        return stringWidth(texto, fuente, _TAMANO_CELDA)

    minimos, deseados = [], []
    for i, columna in enumerate(columnas):
        celdas = [f[i] for f in filas if i < len(f)]
        palabras = [ancho(p, True) for p in columna.split()] + [
            ancho(p) for t in celdas for p in t.split()
        ]
        minimo = max(palabras, default=0) + relleno
        completo = max([ancho(columna, True)] + [ancho(t) for t in celdas]) + relleno
        minimos.append(minimo)
        deseados.append(max(min(completo, ancho_util * 0.35), minimo))

    if sum(minimos) >= ancho_util:
        return [ancho_util * m / sum(minimos) for m in minimos]
    sobrante = ancho_util - sum(minimos)
    extras = [d - m for d, m in zip(deseados, minimos)]
    total_extra = sum(extras)
    if total_extra == 0:
        return [m + sobrante / len(minimos) for m in minimos]
    return [m + sobrante * e / total_extra for m, e in zip(minimos, extras)]
